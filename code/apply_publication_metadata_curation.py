#!/usr/bin/env python3
"""Apply human edits from the publication-curation workbook to resolved metadata.

The visible ``Publication Curation`` sheet is compared with the hidden ``Baseline``
sheet. Only fields a reviewer actually changed are applied. A single curated work can
map to several annual-report records; edits are propagated to every listed record ID.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


SHEET_NAME = "Publication Curation"
BASELINE_SHEET_NAME = "Baseline - Do Not Edit"
KEY_COLUMN = "work_key"
RECORD_IDS_COLUMN = "source_record_ids"
EDITABLE_FIELDS = {
    "title": "canonical_title",
    "publication_year": "publication_year",
    "doi": "doi",
    "authors": "authors",
    "landing_url": "landing_url",
    "abstract": "abstract",
}


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    rendered = str(value).strip()
    return rendered[1:] if rendered.startswith("'=") else rendered


def sheet_records(sheet) -> dict[str, dict[str, str]]:
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [text(value) for value in next(rows)]
    except StopIteration:
        return {}
    required = {KEY_COLUMN, RECORD_IDS_COLUMN, *EDITABLE_FIELDS}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"{sheet.title} is missing columns: {', '.join(missing)}")
    records: dict[str, dict[str, str]] = {}
    for values in rows:
        row = {header: text(value) for header, value in zip(headers, values)}
        if row[KEY_COLUMN]:
            records[row[KEY_COLUMN]] = row
    return records


def changed_values(current: dict[str, str], baseline: dict[str, str]) -> dict[str, str]:
    return {
        output_field: current[column]
        for column, output_field in EDITABLE_FIELDS.items()
        if current[column] != baseline[column]
    }


def apply_curation(workbook_path: Path, resolution_path: Path) -> dict[str, int]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if SHEET_NAME not in workbook.sheetnames or BASELINE_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Workbook must contain {SHEET_NAME!r} and {BASELINE_SHEET_NAME!r}")
    current = sheet_records(workbook[SHEET_NAME])
    baseline = sheet_records(workbook[BASELINE_SHEET_NAME])

    edits_by_record: dict[str, dict[str, str]] = {}
    changed_works = 0
    changed_fields = 0
    for work_key, row in current.items():
        if work_key not in baseline:
            raise ValueError(f"No baseline row found for {work_key}")
        changes = changed_values(row, baseline[work_key])
        if not changes:
            continue
        changed_works += 1
        changed_fields += len(changes)
        record_ids = [value.strip() for value in row[RECORD_IDS_COLUMN].split(";") if value.strip()]
        for record_id in record_ids:
            edits_by_record.setdefault(record_id, {}).update(changes)

    with resolution_path.open(encoding="utf-8", newline="") as handle:
        reader = csv.DictReader(handle)
        fields = list(reader.fieldnames or [])
        rows = list(reader)

    applied_records = 0
    found_ids: set[str] = set()
    for row in rows:
        changes = edits_by_record.get(row.get("record_id", ""))
        if not changes:
            continue
        found_ids.add(row["record_id"])
        row.update(changes)
        if "abstract" in changes:
            row["abstract_source"] = "human_curated" if changes["abstract"] else ""
            row["local_abstract"] = ""
        if any(field in changes for field in ("canonical_title", "publication_year", "doi", "authors", "landing_url")):
            row["match_source"] = "human_curated"
        applied_records += 1

    missing_ids = sorted(set(edits_by_record) - found_ids)
    if missing_ids:
        sample = ", ".join(missing_ids[:5])
        raise ValueError(f"Curated rows reference {len(missing_ids)} missing record IDs (for example: {sample})")

    temporary = resolution_path.with_suffix(resolution_path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(resolution_path)
    return {
        "curated_works": changed_works,
        "curated_fields": changed_fields,
        "updated_resolution_records": applied_records,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=Path("data/publication_metadata_curation.xlsx"))
    parser.add_argument("--resolution", type=Path, default=Path("results/publication_resolution/publication_resolution.csv"))
    args = parser.parse_args()
    if not args.workbook.exists():
        print(f"Curation workbook not found; no human edits applied: {args.workbook}")
        return
    summary = apply_curation(args.workbook, args.resolution)
    print("Applied publication metadata curation: " + ", ".join(f"{key}={value}" for key, value in summary.items()))


if __name__ == "__main__":
    main()
