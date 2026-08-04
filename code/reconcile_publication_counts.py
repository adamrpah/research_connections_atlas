#!/usr/bin/env python3
"""Compare extracted annual-report publication candidates with workbook totals."""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


def normalized_title(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()


def workbook_counts(path: Path) -> dict[int, int]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    headers = {str(cell.value).strip(): i for i, cell in enumerate(sheet[1], start=1)}
    year_col = headers["Year"]
    count_col = headers["Scholarly Papers Published"]
    return {
        int(row[year_col - 1].value): int(row[count_col - 1].value)
        for row in sheet.iter_rows(min_row=2)
        if row[year_col - 1].value is not None and row[count_col - 1].value is not None
    }


def extracted_counts(path: Path) -> dict[int, dict]:
    grouped: dict[int, list[dict]] = defaultdict(list)
    with path.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            if row["report_year"]:
                grouped[int(row["report_year"])].append(row)

    output = {}
    for year, rows in grouped.items():
        titled = [r for r in rows if r["title"].strip()]
        high_conf = [r for r in titled if float(r["confidence"]) >= 0.75]
        unique_titles = {normalized_title(r["title"]) for r in titled if normalized_title(r["title"])}
        output[year] = {
            "candidate_count": len(rows),
            "parsed_title_count": len(titled),
            "high_confidence_title_count": len(high_conf),
            "unique_normalized_title_count": len(unique_titles),
            "review_count": sum(r["needs_review"].lower() == "true" for r in rows),
        }
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=Path("data/Annual Report compilation.xlsx"))
    parser.add_argument("--candidates", type=Path, default=Path("results/publication_candidates.csv"))
    parser.add_argument("--output", type=Path, default=Path("results/publication_count_reconciliation.csv"))
    parser.add_argument("--summary", type=Path, default=Path("results/publication_count_reconciliation.json"))
    parser.add_argument("--exclude-year", type=int, action="append", default=[1999],
                        help="Year to omit from reconciliation (repeatable; default: 1999)")
    args = parser.parse_args()

    excluded = set(args.exclude_year)
    expected = {year: count for year, count in workbook_counts(args.workbook).items() if year not in excluded}
    extracted = extracted_counts(args.candidates)
    extracted = {year: stats for year, stats in extracted.items() if year not in excluded}
    rows = []
    for year in sorted(set(expected) | set(extracted)):
        listed = expected.get(year)
        stats = extracted.get(year, {})
        candidates = stats.get("candidate_count", 0)
        parsed = stats.get("parsed_title_count", 0)
        delta = candidates - listed if listed is not None else None
        rows.append({
            "year": year,
            "workbook_publications": listed,
            **{key: stats.get(key, 0) for key in (
                "candidate_count", "parsed_title_count", "high_confidence_title_count",
                "unique_normalized_title_count", "review_count"
            )},
            "candidate_delta": delta,
            "candidate_ratio": round(candidates / listed, 4) if listed else None,
            "absolute_candidate_delta": abs(delta) if delta is not None else None,
            "status": (
                "not_in_workbook" if listed is None else
                "no_extraction" if candidates == 0 else
                "exact" if delta == 0 else
                "within_5_percent" if abs(delta) / listed <= 0.05 else
                "over" if delta > 0 else "under"
            ),
        })

    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)

    comparable = [r for r in rows if r["workbook_publications"] is not None]
    summary = {
        "source_workbook": str(args.workbook),
        "source_candidates": str(args.candidates),
        "excluded_years": sorted(excluded),
        "years_in_workbook": len(expected),
        "workbook_total": sum(expected.values()),
        "candidate_total_for_workbook_years": sum(r["candidate_count"] for r in comparable),
        "parsed_title_total_for_workbook_years": sum(r["parsed_title_count"] for r in comparable),
        "years_exact_or_within_5_percent": sum(r["status"] in {"exact", "within_5_percent"} for r in comparable),
        "years_over": sum(r["status"] == "over" for r in comparable),
        "years_under": sum(r["status"] == "under" for r in comparable),
        "years_without_extraction": [r["year"] for r in comparable if r["status"] == "no_extraction"],
        "rows": rows,
    }
    args.summary.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in summary.items() if key != "rows"}, indent=2))


if __name__ == "__main__":
    main()
