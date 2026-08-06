import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from import_faculty_attribution_review import (
    derive_attribution_overrides, derive_name_overrides, parse_review_workbook,
)


class FacultyAttributionReviewImportTests(unittest.TestCase):
    def setUp(self):
        self.registry = [
            {"faculty_id": "fac_jane", "faculty_name": "Jane Scholar"},
            {"faculty_id": "fac_john", "faculty_name": "John Writer"},
        ]

    def make_workbook(self, path: Path):
        workbook = Workbook()
        canonical = workbook.active
        canonical.title = "Canonical Faculty"
        canonical.append([
            "review_action", "canonical_name_override", "alias_to_add", "reviewer",
            "reviewed_at", "reviewer_notes", "faculty_id", "faculty_name",
        ])
        canonical.append([
            "rename", "Jane Q. Scholar", "", "Reviewer", "2026-08-06",
            "Preferred full name", "fac_jane", "Jane Scholar",
        ])

        attribution = workbook.create_sheet("Attribution Review")
        attribution.append([
            "review_action", "replacement_faculty_name", "reviewer", "reviewed_at",
            "reviewer_notes", "publication_id", "faculty_name",
        ])
        attribution.append([
            "replace", "John Writer", "Reviewer", "2026-08-06", "Corrected author",
            "pub_1", "Jane Scholar",
        ])

        exceptions = workbook.create_sheet("Exceptions")
        exceptions.append([
            "review_action", "replacement_faculty_name", "reviewer", "reviewed_at",
            "reviewer_notes", "publication_id", "proposed_faculty_name",
        ])
        exceptions.append([
            "approve_add", "", "Reviewer", "2026-08-06", "Citation confirms author",
            "pub_2", "Jane Scholar",
        ])
        workbook.save(path)

    def test_imports_decisions_and_derives_overrides(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "review.xlsx"
            self.make_workbook(path)
            decisions = parse_review_workbook(path, self.registry)

        self.assertEqual(len(decisions), 3)
        name_overrides = derive_name_overrides(decisions)
        self.assertEqual(name_overrides[0]["action"], "rename")
        self.assertEqual(name_overrides[0]["value"], "Jane Q. Scholar")

        attribution_overrides = derive_attribution_overrides(decisions)
        self.assertEqual(
            [(row["publication_id"], row["action"], row["faculty_name"])
             for row in attribution_overrides],
            [
                ("pub_1", "remove", "Jane Scholar"),
                ("pub_1", "add", "John Writer"),
                ("pub_2", "add", "Jane Scholar"),
            ],
        )


if __name__ == "__main__":
    unittest.main()
