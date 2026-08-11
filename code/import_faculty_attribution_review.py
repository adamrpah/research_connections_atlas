#!/usr/bin/env python3
"""Import a reviewed faculty-attribution workbook into durable CSV decisions.

The workbook is a human interface, not the source of truth. This importer validates
review actions and writes a normalized decision log plus the name and attribution
override files consumed by subsequent pipeline runs.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from pathlib import Path

from openpyxl import load_workbook


DECISION_FIELDS = [
    "decision_id", "source_sheet", "target_type", "target_id", "action",
    "faculty_name", "replacement_value", "reason", "reviewer", "reviewed_at",
]
NAME_OVERRIDE_FIELDS = [
    "decision_id", "faculty_id", "action", "value", "reason", "reviewer", "reviewed_at", "source",
]
ATTRIBUTION_OVERRIDE_FIELDS = [
    "feedback_id", "decision_id", "publication_id", "action", "faculty_name",
    "reason", "reviewer", "reviewed_at", "source",
]


def clean(value: object) -> str:
    return " ".join(str(value or "").split())


def read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(path)


def worksheet_rows(workbook, sheet_name: str) -> list[dict]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Required sheet missing: {sheet_name}")
    sheet = workbook[sheet_name]
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return []
    headers = [clean(value) for value in values[0]]
    if not all(headers):
        raise ValueError(f"Blank header in {sheet_name}")
    return [
        dict(zip(headers, row)) for row in values[1:]
        if any(clean(value) for value in row)
    ]


def decision_id(source_sheet: str, target_id: str, faculty_name: str) -> str:
    key = f"{source_sheet}|{target_id}|{faculty_name}".casefold()
    return "review_" + hashlib.sha256(key.encode()).hexdigest()[:20]


def make_decision(
    source_sheet: str, target_type: str, target_id: str, action: str,
    faculty_name: str, replacement: str, reason: str, reviewer: str, reviewed_at: str,
) -> dict:
    return {
        "decision_id": decision_id(source_sheet, target_id, faculty_name),
        "source_sheet": source_sheet, "target_type": target_type, "target_id": target_id,
        "action": action, "faculty_name": faculty_name, "replacement_value": replacement,
        "reason": reason, "reviewer": reviewer, "reviewed_at": reviewed_at,
    }


def parse_review_workbook(path: Path, registry: list[dict]) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    faculty_by_id = {clean(row["faculty_id"]): clean(row["faculty_name"]) for row in registry}
    faculty_names = set(faculty_by_id.values())
    decisions = []

    for row in worksheet_rows(workbook, "Canonical Faculty"):
        action = clean(row.get("review_action")).casefold()
        if not action:
            continue
        if action not in {"approve", "rename", "add_alias", "exclude", "clear"}:
            raise ValueError(f"Invalid Canonical Faculty action: {action}")
        faculty_id = clean(row.get("faculty_id"))
        if faculty_id not in faculty_by_id:
            raise ValueError(f"Unknown faculty_id in Canonical Faculty: {faculty_id}")
        replacement = clean(row.get("canonical_name_override") or row.get("alias_to_add"))
        if action in {"rename", "add_alias"} and not replacement:
            raise ValueError(f"{action} requires a replacement value for {faculty_id}")
        decisions.append(make_decision(
            "Canonical Faculty", "faculty", faculty_id, action, faculty_by_id[faculty_id],
            replacement, clean(row.get("reviewer_notes")), clean(row.get("reviewer")),
            clean(row.get("reviewed_at")),
        ))

    for sheet_name in ("Attribution Review", "Exceptions"):
        allowed = ({"approve", "remove", "replace", "clear"} if sheet_name == "Attribution Review"
                   else {"approve_add", "reject", "replace", "clear"})
        for row in worksheet_rows(workbook, sheet_name):
            action = clean(row.get("review_action")).casefold()
            if not action:
                continue
            if action not in allowed:
                raise ValueError(f"Invalid {sheet_name} action: {action}")
            publication_id = clean(row.get("publication_id"))
            faculty_name = clean(row.get("faculty_name") or row.get("proposed_faculty_name"))
            replacement = clean(row.get("replacement_faculty_name"))
            if action == "replace" and replacement not in faculty_names:
                raise ValueError(f"Replacement faculty is not in the registry: {replacement}")
            if action in {"remove", "approve_add"} and faculty_name not in faculty_names:
                raise ValueError(f"Reviewed faculty is not in the registry: {faculty_name}")
            if not publication_id:
                raise ValueError(f"Missing publication_id in {sheet_name}")
            decisions.append(make_decision(
                sheet_name, "publication_attribution", publication_id, action, faculty_name,
                replacement, clean(row.get("reviewer_notes")), clean(row.get("reviewer")),
                clean(row.get("reviewed_at")),
            ))
    return decisions


def derive_name_overrides(decisions: list[dict]) -> list[dict]:
    return [{
        "decision_id": row["decision_id"], "faculty_id": row["target_id"],
        "action": row["action"], "value": row["replacement_value"], "reason": row["reason"],
        "reviewer": row["reviewer"], "reviewed_at": row["reviewed_at"], "source": "review_workbook",
    } for row in decisions
      if row["source_sheet"] == "Canonical Faculty" and row["action"] not in {"approve", "clear"}]


def derive_attribution_overrides(decisions: list[dict]) -> list[dict]:
    rows = []
    for decision in decisions:
        if decision["target_type"] != "publication_attribution":
            continue
        common = {
            "feedback_id": "", "decision_id": decision["decision_id"],
            "publication_id": decision["target_id"], "reason": decision["reason"],
            "reviewer": decision["reviewer"], "reviewed_at": decision["reviewed_at"],
            "source": "review_workbook",
        }
        if decision["action"] == "remove":
            rows.append({**common, "action": "remove", "faculty_name": decision["faculty_name"]})
        elif decision["action"] == "approve_add":
            rows.append({**common, "action": "add", "faculty_name": decision["faculty_name"]})
        elif decision["action"] == "replace":
            if decision["faculty_name"]:
                rows.append({**common, "action": "remove", "faculty_name": decision["faculty_name"]})
            rows.append({**common, "action": "add", "faculty_name": decision["replacement_value"]})
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=Path("data/faculty_attribution_review.xlsx"))
    parser.add_argument("--registry", type=Path, default=Path("results/institutional_faculty_registry.csv"))
    parser.add_argument("--decisions", type=Path, default=Path("data/faculty_attribution_review_decisions.csv"))
    parser.add_argument("--name-overrides", type=Path, default=Path("data/faculty_name_overrides.csv"))
    parser.add_argument("--attribution-overrides", type=Path, default=Path("data/faculty_attribution_overrides.csv"))
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    registry = read_csv(args.registry)
    imported = parse_review_workbook(args.workbook, registry)
    existing = {row["decision_id"]: row for row in read_csv(args.decisions) if row.get("decision_id")}
    for row in imported:
        if row["action"] == "clear":
            existing.pop(row["decision_id"], None)
        else:
            existing[row["decision_id"]] = row
    decisions = sorted(existing.values(), key=lambda row: (row["source_sheet"], row["target_id"], row["faculty_name"]))

    legacy = [row for row in read_csv(args.attribution_overrides) if row.get("source") != "review_workbook"]
    attribution_overrides = legacy + derive_attribution_overrides(decisions)
    if not args.dry_run:
        write_csv(args.decisions, DECISION_FIELDS, decisions)
        write_csv(args.name_overrides, NAME_OVERRIDE_FIELDS, derive_name_overrides(decisions))
        write_csv(args.attribution_overrides, ATTRIBUTION_OVERRIDE_FIELDS, attribution_overrides)
    print({
        "imported_decisions": len(imported), "active_decisions": len(decisions),
        "name_overrides": len(derive_name_overrides(decisions)),
        "attribution_overrides": len(attribution_overrides), "dry_run": args.dry_run,
    })


if __name__ == "__main__":
    main()
